import csv
import io
import json
from typing import Tuple, Dict, Any, List, Optional
from backend.app.database.store import db_store
from backend.app.models.entities import TimetableSession, Subject, Teacher, ClassBatch, Room, WorkingSettings

class IOService:
    @staticmethod
    def export_timetable_csv() -> str:
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow([
            "Session ID", "Class ID", "Class Name", "Subject Code",
            "Subject Name", "Teacher ID", "Teacher Name", "Room ID",
            "Room Name", "Day", "Period", "Duration", "Status"
        ])

        for sess in db_store.sessions.values():
            cls = db_store.classes.get(sess.class_id)
            sub = db_store.subjects.get(sess.subject_id)
            teacher = db_store.teachers.get(sess.teacher_id)
            room = db_store.rooms.get(sess.room_id)

            writer.writerow([
                sess.id,
                sess.class_id,
                cls.name if cls else sess.class_id,
                sub.code if sub else sess.subject_id,
                sub.name if sub else sess.subject_id,
                sess.teacher_id,
                teacher.name if teacher else sess.teacher_id,
                sess.room_id,
                room.name if room else sess.room_id,
                sess.day,
                sess.period,
                sess.duration,
                sess.status
            ])

        return output.getvalue()

    @staticmethod
    def export_timetable_excel() -> bytes:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

        wb = openpyxl.Workbook()
        # Default sheet -> All Sessions
        ws_all = wb.active
        ws_all.title = "All Sessions"

        header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="3B5998", end_color="3B5998", fill_type="solid")
        border_thin = Border(
            left=Side(style="thin", color="CCCCCC"),
            right=Side(style="thin", color="CCCCCC"),
            top=Side(style="thin", color="CCCCCC"),
            bottom=Side(style="thin", color="CCCCCC")
        )

        headers = [
            "Session ID", "Class Name", "Subject Code", "Subject Name",
            "Teacher Name", "Room Name", "Day", "Period", "Duration", "Status"
        ]
        ws_all.append(headers)
        for col_idx in range(1, len(headers) + 1):
            cell = ws_all.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for sess in db_store.sessions.values():
            cls = db_store.classes.get(sess.class_id)
            sub = db_store.subjects.get(sess.subject_id)
            teacher = db_store.teachers.get(sess.teacher_id)
            room = db_store.rooms.get(sess.room_id)
            row_data = [
                sess.id,
                cls.name if cls else sess.class_id,
                sub.code if sub else sess.subject_id,
                sub.name if sub else sess.subject_id,
                teacher.name if teacher else sess.teacher_id,
                room.name if room else sess.room_id,
                sess.day,
                sess.period,
                sess.duration,
                sess.status
            ]
            ws_all.append(row_data)

        # Classes Matrices
        days = db_store.settings.days
        periods = list(range(1, db_store.settings.periods_per_day + 1))
        
        for c_id, cls in db_store.classes.items():
            ws_cls = wb.create_sheet(title=f"Class_{cls.name[:25]}")
            ws_cls.append([f"Timetable for {cls.name} (Strength: {cls.strength})"])
            ws_cls.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(periods) + 1)
            ws_cls.cell(row=1, column=1).font = Font(name="Arial", size=13, bold=True, color="2C3E50")
            
            grid_headers = ["Day"] + [f"Period {p}" for p in periods]
            ws_cls.append(grid_headers)
            for c_i in range(1, len(grid_headers) + 1):
                c = ws_cls.cell(row=2, column=c_i)
                c.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
                c.fill = PatternFill(start_color="4A69BD", end_color="4A69BD", fill_type="solid")
                c.alignment = Alignment(horizontal="center", vertical="center")

            # Build mapping
            c_sessions = [s for s in db_store.sessions.values() if s.class_id == c_id and s.status != "cancelled"]
            slot_map = {}
            for s in c_sessions:
                for p_off in range(s.duration):
                    slot_map[(s.day, s.period + p_off)] = s

            for d in days:
                row_vals = [d]
                for p in periods:
                    if p == cls.lunch_break_period:
                        row_vals.append("LUNCH BREAK")
                    elif (d, p) in slot_map:
                        sess = slot_map[(d, p)]
                        sub = db_store.subjects.get(sess.subject_id)
                        tea = db_store.teachers.get(sess.teacher_id)
                        rm = db_store.rooms.get(sess.room_id)
                        txt = f"{sub.code if sub else sess.subject_id}\n{tea.name if tea else sess.teacher_id}\n{rm.name if rm else sess.room_id}"
                        row_vals.append(txt)
                    else:
                        row_vals.append("-")
                ws_cls.append(row_vals)

        # Teachers Matrices
        for t_id, tea in db_store.teachers.items():
            ws_tea = wb.create_sheet(title=f"Teacher_{tea.name[:23]}")
            ws_tea.append([f"Timetable for {tea.name} (Max Load: {tea.max_weekly_load}h)"])
            ws_tea.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(periods) + 1)
            ws_tea.cell(row=1, column=1).font = Font(name="Arial", size=13, bold=True, color="2C3E50")
            
            grid_headers = ["Day"] + [f"Period {p}" for p in periods]
            ws_tea.append(grid_headers)
            for c_i in range(1, len(grid_headers) + 1):
                c = ws_tea.cell(row=2, column=c_i)
                c.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
                c.fill = PatternFill(start_color="1E3799", end_color="1E3799", fill_type="solid")
                c.alignment = Alignment(horizontal="center", vertical="center")

            t_sessions = [s for s in db_store.sessions.values() if s.teacher_id == t_id and s.status != "cancelled"]
            slot_map = {}
            for s in t_sessions:
                for p_off in range(s.duration):
                    slot_map[(s.day, s.period + p_off)] = s

            for d in days:
                row_vals = [d]
                for p in periods:
                    if (d, p) in slot_map:
                        sess = slot_map[(d, p)]
                        cls = db_store.classes.get(sess.class_id)
                        sub = db_store.subjects.get(sess.subject_id)
                        rm = db_store.rooms.get(sess.room_id)
                        txt = f"{cls.name if cls else sess.class_id}\n{sub.code if sub else sess.subject_id}\n{rm.name if rm else sess.room_id}"
                        row_vals.append(txt)
                    else:
                        row_vals.append("-")
                ws_tea.append(row_vals)

        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()

    @staticmethod
    def export_timetable_pdf() -> bytes:
        from reportlab.lib.pagesizes import letter, landscape
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib import colors

        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=landscape(letter),
            rightMargin=20,
            leftMargin=20,
            topMargin=20,
            bottomMargin=20
        )
        elements = []
        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            name="TitleStyle",
            parent=styles["Heading1"],
            fontSize=18,
            textColor=colors.HexColor("#2C3E50"),
            spaceAfter=10
        )
        sub_style = ParagraphStyle(
            name="SubStyle",
            parent=styles["Normal"],
            fontSize=10,
            textColor=colors.HexColor("#7F8C8D"),
            spaceAfter=15
        )
        cell_style = ParagraphStyle(
            name="CellStyle",
            parent=styles["Normal"],
            fontSize=8,
            leading=10,
            alignment=1
        )

        elements.append(Paragraph("SchedulAI — Official Institutional Timetable", title_style))
        elements.append(Paragraph(f"Generated schedule with {len(db_store.sessions)} active sessions.", sub_style))

        days = db_store.settings.days
        periods = list(range(1, db_store.settings.periods_per_day + 1))

        for c_id, cls in db_store.classes.items():
            elements.append(Spacer(1, 10))
            elements.append(Paragraph(f"<b>Cohort: {cls.name} (Strength: {cls.strength})</b>", styles["Heading3"]))
            
            c_sessions = [s for s in db_store.sessions.values() if s.class_id == c_id and s.status != "cancelled"]
            slot_map = {}
            for s in c_sessions:
                for p_off in range(s.duration):
                    slot_map[(s.day, s.period + p_off)] = s

            headers = ["Day"] + [f"Period {p}" for p in periods]
            table_data = [headers]

            for d in days:
                row = [d]
                for p in periods:
                    if p == cls.lunch_break_period:
                        row.append(Paragraph("<b>LUNCH BREAK</b>", cell_style))
                    elif (d, p) in slot_map:
                        sess = slot_map[(d, p)]
                        sub = db_store.subjects.get(sess.subject_id)
                        tea = db_store.teachers.get(sess.teacher_id)
                        rm = db_store.rooms.get(sess.room_id)
                        tag = "[LAB]" if sub and sub.type == "lab" else ""
                        text = f"<b>{sub.code if sub else sess.subject_id} {tag}</b><br/>{tea.name if tea else sess.teacher_id}<br/>{rm.name if rm else sess.room_id}"
                        row.append(Paragraph(text, cell_style))
                    else:
                        row.append(Paragraph("-", cell_style))
                table_data.append(row)

            col_widths = [65] + [100] * len(periods)
            t = Table(table_data, colWidths=col_widths, repeatRows=1)
            t.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#34495E")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#BDC3C7")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.HexColor("#FAFAFA"), colors.HexColor("#F0F3F4")]),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]))
            elements.append(t)
            elements.append(Spacer(1, 15))

        doc.build(elements)
        return buffer.getvalue()

    @staticmethod
    def import_dataset_json(data: dict) -> Tuple[bool, str]:
        try:
            for key in ["subjects", "teachers", "classes", "rooms"]:
                if key not in data or not isinstance(data[key], list):
                    return False, f"Missing or invalid section: '{key}'"

            if len(data["classes"]) == 0 or len(data["teachers"]) == 0 or len(data["rooms"]) == 0:
                return False, "Dataset must contain at least one class, teacher, and room"

            db_store.load_from_dict(data)
            return True, f"Successfully imported {len(db_store.classes)} classes, {len(db_store.teachers)} teachers, {len(db_store.rooms)} rooms."
        except Exception as e:
            return False, f"Failed to parse and import dataset: {str(e)}"

    @staticmethod
    def import_csv_data(
        subjects_csv: Optional[str] = None,
        teachers_csv: Optional[str] = None,
        classes_csv: Optional[str] = None,
        rooms_csv: Optional[str] = None
    ) -> Tuple[bool, str]:
        """Import entities from CSV strings"""
        try:
            if subjects_csv:
                reader = csv.DictReader(io.StringIO(subjects_csv))
                for row in reader:
                    sub_id = row.get("id") or row.get("code")
                    if sub_id:
                        db_store.subjects[sub_id] = Subject(
                            id=sub_id,
                            code=row.get("code", sub_id),
                            name=row.get("name", sub_id),
                            weekly_hours=int(row.get("weekly_hours", 3)),
                            type=row.get("type", "theory").lower(),
                            lab_duration=int(row.get("lab_duration", 2 if row.get("type") == "lab" else 1))
                        )

            if rooms_csv:
                reader = csv.DictReader(io.StringIO(rooms_csv))
                for row in reader:
                    r_id = row.get("id") or row.get("name")
                    if r_id:
                        db_store.rooms[r_id] = Room(
                            id=r_id,
                            name=row.get("name", r_id),
                            capacity=int(row.get("capacity", 40)),
                            type=row.get("type", "classroom").lower(),
                            active=True
                        )

            if teachers_csv:
                reader = csv.DictReader(io.StringIO(teachers_csv))
                for row in reader:
                    t_id = row.get("id") or row.get("name")
                    sub_ids = [s.strip() for s in row.get("subject_ids", "").split(";") if s.strip()]
                    if t_id:
                        db_store.teachers[t_id] = Teacher(
                            id=t_id,
                            name=row.get("name", t_id),
                            subject_ids=sub_ids,
                            max_weekly_load=int(row.get("max_weekly_load", 15)),
                            unavailable_slots=[],
                            preferred_slots=[],
                            active=True
                        )

            if classes_csv:
                reader = csv.DictReader(io.StringIO(classes_csv))
                for row in reader:
                    c_id = row.get("id") or row.get("name")
                    reqs = {}
                    if row.get("subject_requirements"):
                        pairs = row.get("subject_requirements").split(";")
                        for p in pairs:
                            if ":" in p:
                                k, v = p.split(":")
                                reqs[k.strip()] = int(v.strip())
                    if c_id:
                        db_store.classes[c_id] = ClassBatch(
                            id=c_id,
                            name=row.get("name", c_id),
                            strength=int(row.get("strength", 30)),
                            subject_requirements=reqs,
                            lunch_break_period=int(row.get("lunch_break_period", 4)),
                            active=True
                        )

            return True, f"Successfully parsed CSVs: {len(db_store.classes)} classes, {len(db_store.teachers)} teachers, {len(db_store.rooms)} rooms."
        except Exception as e:
            return False, f"CSV parsing error: {str(e)}"
